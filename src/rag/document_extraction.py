"""Unified text + embedded-image extraction for the single document upload
field (PDF or Word). Extracted images are validated the same way as any
other untrusted upload (real content sniffing, size caps — see
src/security/file_validation.py) before being handed to the vision model.
"""
import io
from pathlib import Path
from typing import List, Optional

from docx import Document
from pypdf import PdfReader

from src.security.file_validation import sniff_image_mime, MAX_IMAGES_PER_REQUEST, MAX_IMAGE_SIZE_BYTES


def _get_extension(source, filename: Optional[str] = None) -> str:
    name = filename or getattr(source, "name", None) or str(source)
    return Path(name).suffix.lower()


def read_pdf_text(source) -> str:
    reader = PdfReader(str(source) if isinstance(source, Path) else source)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def read_docx_text(source) -> str:
    document = Document(str(source) if isinstance(source, Path) else source)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def read_document_text(source, filename: Optional[str] = None) -> str:
    ext = _get_extension(source, filename)
    if hasattr(source, "seek"):
        source.seek(0)
    if ext == ".docx":
        return read_docx_text(source)
    return read_pdf_text(source)


def _extract_images_from_pdf(source) -> List[bytes]:
    reader = PdfReader(str(source) if isinstance(source, Path) else source)
    images = []
    for page in reader.pages:
        for image_file in page.images:
            images.append(image_file.data)
            if len(images) >= MAX_IMAGES_PER_REQUEST * 2:  # collect a small surplus, filtered below
                return images
    return images


def _extract_images_from_docx(source) -> List[bytes]:
    document = Document(str(source) if isinstance(source, Path) else source)
    images = []
    for rel in document.part.rels.values():
        if "image" in rel.reltype:
            try:
                images.append(rel.target_part.blob)
            except (KeyError, AttributeError):
                continue
    return images


def extract_embedded_images(source, filename: Optional[str] = None) -> List[dict]:
    """Return validated embedded images as [{"bytes":..., "mime":..., "name":...}],
    capped at MAX_IMAGES_PER_REQUEST and each within MAX_IMAGE_SIZE_BYTES.
    Any blob that isn't a real, sniffable PNG/JPEG is silently skipped."""
    ext = _get_extension(source, filename)
    if hasattr(source, "seek"):
        source.seek(0)
    try:
        raw_images = _extract_images_from_docx(source) if ext == ".docx" else _extract_images_from_pdf(source)
    except Exception:
        return []

    results = []
    for index, data in enumerate(raw_images):
        if not data or len(data) > MAX_IMAGE_SIZE_BYTES:
            continue
        mime = sniff_image_mime(data)
        if mime is None:
            continue
        results.append({"bytes": data, "mime": mime, "name": f"صورة مضمّنة {index + 1}"})
        if len(results) >= MAX_IMAGES_PER_REQUEST:
            break
    return results
